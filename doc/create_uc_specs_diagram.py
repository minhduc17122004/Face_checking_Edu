# -*- coding: utf-8 -*-
"""Đặc tả UC theo sơ đồ Use Case - Hệ thống điểm danh học sinh → xuất .docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT = r"c:\Users\ADMIN\Desktop\MIS\hrm-flutter_dev_2.0\hrm-flutter_dev_2.0\face_time_keeping\doc\DacTaUseCase_TheoSoDo.docx"
OUTPUT_FALLBACK = r"c:\Users\ADMIN\Desktop\MIS\hrm-flutter_dev_2.0\hrm-flutter_dev_2.0\face_time_keeping\doc\DacTaUseCase_ChimucDrawio.docx"
DARK_BLUE = RGBColor(0x1F, 0x49, 0x7D)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_cell_bg(cell, hx):
    tc = cell._tc
    p = tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), hx)
    p.append(s)


def set_cell_borders(cell, sz=4):
    tc = cell._tc
    p = tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        x = OxmlElement(f"w:{side}")
        x.set(qn("w:val"), "single")
        x.set(qn("w:sz"), str(sz))
        x.set(qn("w:space"), "0")
        x.set(qn("w:color"), "2E74B5")
        b.append(x)
    p.append(b)


def set_col_w(cell, cm):
    tc = cell._tc
    p = tc.get_or_add_tcPr()
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(int(cm * 567)))
    w.set(qn("w:type"), "dxa")
    p.append(w)


def cell_txt(cell, text, bold=False, first=True, color=BLACK, size=12):
    if first:
        cell.text = ""
        para = cell.paragraphs[0]
    else:
        para = cell.add_paragraph()
    r = para.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = "Times New Roman"


def tbl_borders(tbl):
    t = tbl._tbl
    pr = t.find(qn("w:tblPr"))
    if pr is None:
        pr = OxmlElement("w:tblPr")
        t.insert(0, pr)
    bs = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{side}")
        x.set(qn("w:val"), "single")
        x.set(qn("w:sz"), "6")
        x.set(qn("w:space"), "0")
        x.set(qn("w:color"), "2E74B5")
        bs.append(x)
    pr.append(bs)


def add_uc_table(doc, title, rows):
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = DARK_BLUE
    r.font.name = "Times New Roman"
    p.paragraph_format.space_before = Cm(0.4)
    p.paragraph_format.space_after = Cm(0.2)

    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_borders(t)
    for i, (lab, val) in enumerate(rows):
        row = t.add_row()
        c0, c1 = row.cells[0], row.cells[1]
        set_col_w(c0, 4.2)
        set_col_w(c1, 11.3)
        set_cell_bg(c0, "D6E4F0" if i % 2 == 0 else "EEF4FA")
        set_cell_bg(c1, "FFFFFF")
        cell_txt(c0, lab, bold=True, color=DARK_BLUE)
        cell_txt(c1, val, bold=False)
        set_cell_borders(c0)
        set_cell_borders(c1)
    doc.add_paragraph()


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(3.0)
        s.right_margin = Cm(2.5)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(13)

    h = doc.add_paragraph()
    hr = h.add_run(
        "ĐẶC TẢ CHI TIẾT USE CASE\n"
        "Hệ thống điểm danh học sinh\n"
        "(Theo chỉ mục UC trên sơ đồ Draw.io)"
    )
    hr.bold = True
    hr.font.size = Pt(14)
    hr.font.name = "Times New Roman"
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph()
    ir = intro.add_run(
        "Tài liệu mô tả đặc tả chi tiết các Use Case theo đúng chỉ mục trên sơ đồ "
        "(UC 1, 1.1, 1.2, 2…13). Mỗi Use Case được trình bày dưới dạng bảng gồm: "
        "mã UC, tác nhân, mô tả, tiền điều kiện, hậu điều kiện, luồng chính, luồng phụ, "
        "quan hệ Include/Extend (nếu có), yêu cầu phi chức năng."
    )
    ir.font.name = "Times New Roman"
    ir.font.size = Pt(13)

    rel = doc.add_paragraph()
    rr = rel.add_run(
        "Tóm tắt quan hệ theo sơ đồ: UC 1 (Quản lý chứng thực) gồm UC 1.1 "
        "(Đăng nhập) và UC 1.2 (Đăng xuất) <<extend>> UC 1. Các UC 2, 3, 4, 5, "
        "6, 7, 8, 9, 10, 11, 12, 13 đều <<include>> UC 1.1. UC 4.1, 4.2, 4.3 "
        "<<extend>> UC 4. UC 5.1, 6.1, 10.1 <<extend>> UC 1.1."
    )
    rr.font.name = "Times New Roman"
    rr.font.size = Pt(13)

    # Load specs from external data file (same folder)
    import json
    import os
    base = os.path.dirname(os.path.abspath(__file__))
    jpath = os.path.join(base, "uc_specs_diagram_data.json")
    with open(jpath, "r", encoding="utf-8") as f:
        specs = json.load(f)

    for item in specs:
        add_uc_table(doc, item["heading"], item["rows"])

    try:
        doc.save(OUTPUT)
        print("OK:", OUTPUT)
    except PermissionError:
        doc.save(OUTPUT_FALLBACK)
        # ASCII-only: Windows console may not be UTF-8.
        print("WARN: primary locked, saved:", OUTPUT_FALLBACK)


if __name__ == "__main__":
    main()
